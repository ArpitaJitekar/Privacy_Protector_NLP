"""
Document text extraction module supporting PDF, DOCX, and TXT files.
Preserves location metadata (page, paragraph, or line numbers) for detailed reporting.
"""

from dataclasses import dataclass
from typing import List, Optional
import io


@dataclass
class TextSegment:
    """Represents a discrete block of text with location metadata."""
    text: str
    location: str          # e.g., "Page 1, Line 4", "Paragraph 3", "Line 12"
    page: Optional[int] = None
    line: Optional[int] = None
    section: Optional[str] = None


class DocumentExtractor:
    """Extracts text segments from various document formats."""

    @staticmethod
    def extract_from_txt(content_bytes: bytes, filename: str = "document.txt") -> List[TextSegment]:
        """Extracts text from plain text bytes."""
        segments: List[TextSegment] = []
        try:
            text = content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content_bytes.decode("latin-1")
            except Exception:
                text = content_bytes.decode("utf-8", errors="replace")

        lines = text.splitlines()
        for idx, line in enumerate(lines, start=1):
            line_str = line.strip()
            if line_str:
                segments.append(
                    TextSegment(
                        text=line_str,
                        location=f"Line {idx}",
                        line=idx
                    )
                )
        return segments

    @staticmethod
    def extract_from_docx(content_bytes: bytes, filename: str = "document.docx") -> List[TextSegment]:
        """Extracts text paragraphs and table cells from DOCX files using python-docx."""
        import docx

        segments: List[TextSegment] = []
        doc_stream = io.BytesIO(content_bytes)
        doc = docx.Document(doc_stream)

        # Paragraphs
        for idx, p in enumerate(doc.paragraphs, start=1):
            text = p.text.strip()
            if text:
                segments.append(
                    TextSegment(
                        text=text,
                        location=f"Paragraph {idx}",
                        line=idx,
                        section="Paragraph"
                    )
                )

        # Tables
        for table_idx, table in enumerate(doc.tables, start=1):
            for row_idx, row in enumerate(table.rows, start=1):
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    joined = " | ".join(dict.fromkeys(row_texts))  # remove dupes from merged cells
                    segments.append(
                        TextSegment(
                            text=joined,
                            location=f"Table {table_idx}, Row {row_idx}",
                            line=row_idx,
                            section="Table"
                        )
                    )
        return segments

    @staticmethod
    def extract_from_pdf(content_bytes: bytes, filename: str = "document.pdf") -> List[TextSegment]:
        """Extracts text from PDF documents using PyMuPDF (fitz) with page numbers."""
        try:
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz
        except ImportError:
            raise ImportError(
                "PyMuPDF is required to process PDF files. Please install it using 'pip install pymupdf'."
            )

        segments: List[TextSegment] = []
        doc = fitz.open(stream=content_bytes, filetype="pdf")

        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            lines = text.splitlines()
            for line_idx, line in enumerate(lines, start=1):
                clean_line = line.strip()
                if clean_line:
                    segments.append(
                        TextSegment(
                            text=clean_line,
                            location=f"Page {page_num + 1}, Line {line_idx}",
                            page=page_num + 1,
                            line=line_idx
                        )
                    )
        doc.close()
        return segments

    @classmethod
    def extract(cls, file_bytes: bytes, filename: str) -> List[TextSegment]:
        """Auto-detects document format by extension and extracts structured segments."""
        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            return cls.extract_from_pdf(file_bytes, filename)
        elif lower_name.endswith(".docx"):
            return cls.extract_from_docx(file_bytes, filename)
        elif lower_name.endswith((".txt", ".log", ".csv", ".tsv", ".md")):
            return cls.extract_from_txt(file_bytes, filename)
        else:
            # Fallback to plain text decoding
            return cls.extract_from_txt(file_bytes, filename)
